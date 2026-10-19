import os
import tempfile
from pathlib import Path
import pytest
import docx
import fitz

from backend.app.services.document_loader import (
    safe_filename,
    extract_text_from_txt,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_document_text,
)


def test_safe_filename():
    assert safe_filename("../../test.pdf") == "test.pdf"
    assert safe_filename("normal_file.txt") == "normal_file.txt"
    assert safe_filename("weird!@#$name.docx") == "weird____name.docx"


def test_extract_text_from_txt():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("Line 1\nLine 2\x0cPage 2 content")
        temp_path = f.name

    try:
        pages, is_scanned = extract_text_from_txt(temp_path)
        assert len(pages) == 2
        assert "Line 1" in pages[0]["text"]
        assert "Page 2 content" in pages[1]["text"]
        assert not is_scanned
    finally:
        os.remove(temp_path)


def test_extract_text_from_docx():
    doc = docx.Document()
    doc.add_paragraph("Hello DOCX paragraph 1")
    doc.add_paragraph("Paragraph 2 content")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Col1"
    table.rows[0].cells[1].text = "Col2"

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name
    doc.save(temp_path)

    try:
        pages, is_scanned = extract_text_from_docx(temp_path)
        assert len(pages) >= 1
        assert "Hello DOCX" in pages[0]["text"]
        assert "Col1 | Col2" in pages[0]["text"]
        assert not is_scanned
    finally:
        os.remove(temp_path)


def test_extract_text_from_pdf():
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "PDF Page 1 text for testing")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "PDF Page 2 additional information")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name
    doc.save(temp_path)
    doc.close()

    try:
        pages, is_scanned = extract_text_from_pdf(temp_path)
        assert len(pages) == 2
        assert pages[0]["page_number"] == 1
        assert "PDF Page 1 text" in pages[0]["text"]
        assert pages[1]["page_number"] == 2
        assert "Page 2 additional" in pages[1]["text"]
        assert not is_scanned
    finally:
        os.remove(temp_path)


def test_empty_document_detection():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as f:
        temp_path = f.name

    try:
        pages, is_scanned = extract_text_from_txt(temp_path)
        assert len(pages) == 0
        assert is_scanned
    finally:
        os.remove(temp_path)
