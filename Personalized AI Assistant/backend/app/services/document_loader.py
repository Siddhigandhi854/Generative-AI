import os
import re
import uuid
from pathlib import Path
from typing import List, Dict, Any, Tuple

import fitz  # PyMuPDF
import docx

from backend.app.config import STORAGE_DIR


def safe_filename(filename: str) -> str:
    """Sanitize filename to prevent directory traversal or invalid characters."""
    clean = Path(filename).name
    clean = re.sub(r'[^\w\s.-]', '_', clean).strip()
    return clean or "document"


def extract_text_from_pdf(file_path: str) -> Tuple[List[Dict[str, Any]], bool]:
    """
    Extract text per page from a PDF file using PyMuPDF.
    Returns (pages_list, is_scanned).
    """
    doc = fitz.open(file_path)
    pages = []
    total_text_length = 0

    for page_index in range(len(doc)):
        page = doc.load_page(page_index)
        text = page.get_text("text").strip()
        total_text_length += len(text)
        pages.append({
            "page_number": page_index + 1,
            "text": text,
        })
    doc.close()

    # Detect if scanned: pages exist but zero or virtually no selectable text extracted
    is_scanned = len(pages) > 0 and total_text_length < 20
    return pages, is_scanned


def extract_text_from_docx(file_path: str) -> Tuple[List[Dict[str, Any]], bool]:
    """
    Extract text from DOCX file including paragraphs and tables.
    Returns (pages_list, is_scanned).
    """
    document = docx.Document(file_path)
    text_blocks = []

    for para in document.paragraphs:
        stripped = para.text.strip()
        if stripped:
            text_blocks.append(stripped)

    # Extract tables if any
    for table in document.tables:
        table_rows = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                table_rows.append(row_text)
        if table_rows:
            text_blocks.append("\n".join(table_rows))

    full_text = "\n\n".join(text_blocks).strip()
    is_scanned = len(full_text) == 0
    pages = [{"page_number": 1, "text": full_text}] if full_text else []
    return pages, is_scanned


def extract_text_from_txt(file_path: str) -> Tuple[List[Dict[str, Any]], bool]:
    """
    Extract text from TXT file with UTF-8 / fallback encodings.
    Splits on form-feed characters if present, otherwise returns page 1.
    """
    content = ""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(file_path, "r", encoding=encoding, errors="strict") as handle:
                content = handle.read()
                break
        except (UnicodeDecodeError, UnicodeError):
            continue

    if not content:
        with open(file_path, "r", encoding="utf-8", errors="replace") as handle:
            content = handle.read()

    pages = []
    # Split on form feed character if multi-page TXT
    raw_pages = content.split("\x0c") if "\x0c" in content else [content]
    for idx, p_text in enumerate(raw_pages):
        stripped = p_text.strip()
        if stripped:
            pages.append({"page_number": idx + 1, "text": stripped})

    is_scanned = len(pages) == 0
    return pages, is_scanned


def extract_document_text(file_path: str, extension: str) -> Tuple[List[Dict[str, Any]], bool]:
    """
    Universal dispatcher for extracting document text.
    Returns (pages_list, is_scanned_flag).
    """
    ext = extension.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in {".txt", ".text", ".md"}:
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {extension}")


def save_uploaded_file(file, original_filename: str) -> Tuple[str, str]:
    """
    Safely writes an uploaded file to STORAGE_DIR.
    Returns (absolute_file_path, unique_stored_name).
    """
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    clean = safe_filename(original_filename)
    unique_stored_name = f"{uuid.uuid4().hex[:12]}_{clean}"
    destination = STORAGE_DIR / unique_stored_name

    file.file.seek(0)
    with open(destination, "wb") as buffer:
        while chunk := file.file.read(1024 * 64):
            buffer.write(chunk)
    file.file.seek(0)

    return str(destination), unique_stored_name
